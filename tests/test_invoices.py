from datetime import datetime
from docx import Document
from conftest import find_client
from calendar_processor import create_docx, generate_invoice, format_invoice_number


def test_invoice_number_contains_client_id_and_date():
    assert format_invoice_number("CL001", datetime(2026, 10, 28)) == "CL00128102026"


def test_generated_docx_contains_fictional_client_and_total(clients, tmp_path):
    client = find_client(clients, "Emma de Vries")
    output = tmp_path / "invoice.docx"
    generate_invoice(
        client, datetime(2026, 10, 28),
        [("05-10-2026", "1.0", "€ 100.00"), ("12-10-2026", "0.5", "€ 50.00")],
        "€ 150.00", str(output),
    )
    assert output.is_file()
    doc = Document(output)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Emma de Vries" in text
    assert "CL00128102026" in text
    assert "05-10-2026" in text
    assert "12-10-2026" in text
    assert "€ 150.00" in text


def test_custom_invoice_recipient(clients):
    client = find_client(clients, "Sophie Bakker")
    doc = create_docx(client, datetime(2026, 10, 28), [], "€ 0.00")
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Administratie Sophie" in text
